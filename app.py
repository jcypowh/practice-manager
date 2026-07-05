import os
import ast
import json
import operator
import sqlite3
import datetime
from pathlib import Path
from io import BytesIO

from flask import Flask, g, render_template, request, redirect, url_for, flash, jsonify, Response
import openpyxl

import importer
import claude_integration

_SAFE_OPS = {
    ast.Add: operator.add, ast.Sub: operator.sub,
    ast.Mult: operator.mul, ast.Div: operator.truediv,
    ast.USub: operator.neg, ast.UAdd: operator.pos,
}


def _safe_eval_node(node):
    if isinstance(node, ast.Expression):
        return _safe_eval_node(node.body)
    if isinstance(node, ast.BinOp) and type(node.op) in _SAFE_OPS:
        return _SAFE_OPS[type(node.op)](_safe_eval_node(node.left), _safe_eval_node(node.right))
    if isinstance(node, ast.UnaryOp) and type(node.op) in _SAFE_OPS:
        return _SAFE_OPS[type(node.op)](_safe_eval_node(node.operand))
    if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
        return node.value
    raise ValueError('unsupported expression')


def parse_money_input(raw):
    """Accepts a plain number or a simple arithmetic expression, Excel-style ('=' prefix
    optional): '2100', '=2300*0.75', '2200-360'. Returns (value: float, formula: str) -
    formula is the raw text as typed (so it can be shown again next time), value is the
    computed number actually used for earnings/board display."""
    text = (raw or '').strip()
    if not text:
        return 0.0, ''
    expr = text[1:].strip() if text.startswith('=') else text
    try:
        tree = ast.parse(expr, mode='eval')
        value = float(_safe_eval_node(tree))
    except (SyntaxError, ValueError, ZeroDivisionError, TypeError):
        raise ValueError(f'"{text}" is not a valid number or calculation')
    return value, text

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = Path(os.environ.get('STORAGE_DIR', BASE_DIR)) / 'data'
DATA_DIR.mkdir(parents=True, exist_ok=True)
DB_PATH = DATA_DIR / 'practice_manager.db'

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'dev-secret-change-me')
app.config['TEMPLATES_AUTO_RELOAD'] = True

ACTIVITIES = ['Consult', 'Clinic', 'Scope', 'All day', 'Other']
CATEGORIES = ['cdd', 'corporate', 'shore_own', 'scope', 'other']
SCOPE_COLOR = '#2980b9'

SEED_LOCATIONS = [
    # name, category, color
    ('CDD', 'cdd', '#8e44ad'),
    ('Chatswood', 'corporate', '#e67e22'),
    ('Leichhardt', 'corporate', '#e67e22'),
    ('Brookvale', 'corporate', '#e67e22'),
    ('Maroubra', 'corporate', '#e67e22'),
    ('Blacktown', 'corporate', '#e67e22'),
    ('Darlinghurst', 'corporate', '#e67e22'),
    ('Mater', 'shore_own', '#27ae60'),
    ('NBH', 'shore_own', '#27ae60'),
    ('Freshwater', 'scope', '#2980b9'),
    ('ESPH', 'scope', '#2980b9'),
    ('Dee Why', 'scope', '#2980b9'),
]
SCOPE_SITES = ['CDD', 'Mater', 'ESPH', 'Freshwater', 'Dee Why']

# Four top-level groups for the Analysis page, in the display order the user wants.
# CDD is deliberately never split by activity (it's one entity). Mater is the one site
# that genuinely straddles two groups: its Clinic sessions are "SHORE Gastroenterology"
# but its Scope sessions are "Scopes" - so classification must check activity first.
GROUP_ORDER = ['FORHEALTH Medical Centre', 'SHORE Gastroenterology', 'CDD', 'Scopes']
WEEKDAY_ORDER = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday']
ROTATION_LABELS = ['Week 1', 'Week 2', 'Week 3', 'Week 4']


def classify_group(category, activity):
    if category == 'cdd':
        return 'CDD'
    if activity == 'Scope':
        return 'Scopes'
    if category == 'shore_own':
        return 'SHORE Gastroenterology'
    if category == 'corporate':
        return 'FORHEALTH Medical Centre'
    return 'Other'


# Referral pipelines: consult/clinic work at the "source" sites is what generates scope
# bookings at the associated "target" sites. Used both for the on-page ratio and as
# context fed to Claude so it understands this practice's specific referral logic.
REFERRAL_PIPELINES = [
    {
        'label': 'Chatswood + Brookvale + SHORE Gastro → Mater Scope / Freshwater / Dee Why',
        'source_sites': ['Chatswood', 'Brookvale'],
        'source_group': 'SHORE Gastroenterology',
        'target_sites': ['Freshwater', 'Dee Why'],
        'target_site_activity': [('Mater', 'Scope')],
    },
    {
        'label': 'Darlinghurst + Leichhardt + Maroubra + Blacktown → East Sydney Private (ESPH)',
        'source_sites': ['Darlinghurst', 'Leichhardt', 'Maroubra', 'Blacktown'],
        'source_group': None,
        'target_sites': ['ESPH'],
        'target_site_activity': [],
    },
]

# The practice runs a perpetual 4-week rotation (Week 1-4) that is NOT tied to
# calendar week-of-year - it just keeps cycling regardless of month boundaries.
# Anchored against the real source files: Monday 2026-08-03 is confirmed "Week 2".
ROTATION_ANCHOR_MONDAY = datetime.date(2026, 8, 3)
ROTATION_ANCHOR_CODE = 2


def week_rotation_label(monday_date):
    weeks_diff = (monday_date - ROTATION_ANCHOR_MONDAY).days // 7
    code = ((ROTATION_ANCHOR_CODE - 1 + weeks_diff) % 4) + 1
    return f'Week {code}'

# ---------------------------------------------------------------- DB helpers

def get_db():
    if 'db' not in g:
        g.db = sqlite3.connect(str(DB_PATH))
        g.db.row_factory = sqlite3.Row
        g.db.execute('PRAGMA foreign_keys = ON')
    return g.db


@app.teardown_appcontext
def close_db(e=None):
    db = g.pop('db', None)
    if db is not None:
        db.close()


def cfg(key, default=''):
    row = get_db().execute('SELECT value FROM config WHERE key=?', (key,)).fetchone()
    return row['value'] if row else default


def set_cfg(key, value):
    db = get_db()
    db.execute('INSERT OR REPLACE INTO config(key, value) VALUES(?, ?)', (key, value))
    db.commit()


def init_db():
    db = sqlite3.connect(str(DB_PATH))
    db.row_factory = sqlite3.Row
    db.executescript('''
    CREATE TABLE IF NOT EXISTS config (
        key   TEXT PRIMARY KEY,
        value TEXT NOT NULL DEFAULT ''
    );

    CREATE TABLE IF NOT EXISTS locations (
        id                    INTEGER PRIMARY KEY AUTOINCREMENT,
        name                  TEXT NOT NULL UNIQUE,
        category              TEXT DEFAULT 'corporate',
        color                 TEXT DEFAULT '#888888',
        default_am_value      REAL DEFAULT 0,
        default_pm_value      REAL DEFAULT 0,
        default_allday_value  REAL DEFAULT 0,
        active                INTEGER DEFAULT 1,
        sort_order            INTEGER DEFAULT 0,
        notes                 TEXT DEFAULT ''
    );

    CREATE TABLE IF NOT EXISTS location_activity_rates (
        id                    INTEGER PRIMARY KEY AUTOINCREMENT,
        location_id           INTEGER NOT NULL REFERENCES locations(id) ON DELETE CASCADE,
        activity              TEXT NOT NULL,
        color                 TEXT DEFAULT '',
        am_value              REAL DEFAULT 0,
        pm_value              REAL DEFAULT 0,
        allday_value          REAL DEFAULT 0,
        am_value_formula      TEXT DEFAULT '',
        pm_value_formula      TEXT DEFAULT '',
        allday_value_formula  TEXT DEFAULT '',
        UNIQUE(location_id, activity)
    );

    CREATE TABLE IF NOT EXISTS months (
        id            INTEGER PRIMARY KEY AUTOINCREMENT,
        year          INTEGER NOT NULL,
        month         INTEGER NOT NULL,
        label         TEXT NOT NULL,
        status        TEXT DEFAULT 'draft',
        source        TEXT DEFAULT 'import',
        published_at  TEXT,
        archived_at   TEXT,
        created_at    TEXT DEFAULT CURRENT_TIMESTAMP,
        UNIQUE(year, month)
    );

    CREATE TABLE IF NOT EXISTS blocks (
        id              INTEGER PRIMARY KEY AUTOINCREMENT,
        month_id        INTEGER REFERENCES months(id) ON DELETE CASCADE,
        date            TEXT NOT NULL,
        day_name        TEXT DEFAULT '',
        slot            TEXT NOT NULL,
        location_id     INTEGER REFERENCES locations(id),
        activity        TEXT DEFAULT '',
        activity_raw    TEXT DEFAULT '',
        time_note       TEXT DEFAULT '',
        value_override  REAL,
        time_weight     REAL,
        status          TEXT DEFAULT 'scheduled',
        notes           TEXT DEFAULT '',
        source          TEXT DEFAULT 'import',
        created_at      TEXT DEFAULT CURRENT_TIMESTAMP,
        updated_at      TEXT DEFAULT CURRENT_TIMESTAMP
    );
    CREATE INDEX IF NOT EXISTS idx_blocks_date_slot ON blocks(date, slot);

    CREATE TABLE IF NOT EXISTS import_log (
        id                 INTEGER PRIMARY KEY AUTOINCREMENT,
        original_filename  TEXT NOT NULL,
        month_id           INTEGER REFERENCES months(id),
        rows_inserted      INTEGER DEFAULT 0,
        rows_updated       INTEGER DEFAULT 0,
        rows_skipped       INTEGER DEFAULT 0,
        imported_at        TEXT DEFAULT CURRENT_TIMESTAMP
    );

    CREATE TABLE IF NOT EXISTS board_snapshots (
        month_id        INTEGER PRIMARY KEY REFERENCES months(id) ON DELETE CASCADE,
        snapshot_json   TEXT NOT NULL,
        created_at      TEXT DEFAULT CURRENT_TIMESTAMP
    );
    ''')
    db.commit()

    for stmt in [
        "ALTER TABLE location_activity_rates ADD COLUMN am_value_formula TEXT DEFAULT ''",
        "ALTER TABLE location_activity_rates ADD COLUMN pm_value_formula TEXT DEFAULT ''",
        "ALTER TABLE location_activity_rates ADD COLUMN allday_value_formula TEXT DEFAULT ''",
        "ALTER TABLE blocks ADD COLUMN time_weight REAL",
    ]:
        try:
            db.execute(stmt)
        except sqlite3.OperationalError:
            pass
    db.commit()

    defaults = {
        'claude_api_key': '',
        'analysis_trailing_months': '3',
        'utilization_flag_threshold': '1',
        'analysis_overbook_threshold': '10',
    }
    for k, v in defaults.items():
        db.execute('INSERT OR IGNORE INTO config(key, value) VALUES(?, ?)', (k, v))
    db.commit()

    for i, (name, category, color) in enumerate(SEED_LOCATIONS):
        db.execute('''INSERT OR IGNORE INTO locations(name, category, color, sort_order)
                       VALUES(?,?,?,?)''', (name, category, color, i))
    db.commit()

    for name in SCOPE_SITES:
        row = db.execute('SELECT id FROM locations WHERE name=?', (name,)).fetchone()
        if row:
            db.execute('''INSERT OR IGNORE INTO location_activity_rates(location_id, activity, color)
                          VALUES(?, 'Scope', ?)''', (row['id'], SCOPE_COLOR))
    db.commit()
    db.close()


# ---------------------------------------------------------------- resolution helpers

def _effective_color(rate, loc_color):
    if rate is not None and rate['color']:
        return rate['color']
    return loc_color or '#888888'


def _effective_value(value_override, rate, am_default, pm_default, allday_default, slot):
    if value_override is not None:
        return value_override
    if rate is not None:
        val = {'AM': rate['am_value'], 'PM': rate['pm_value']}.get(slot, rate['allday_value'])
    else:
        val = {'AM': am_default, 'PM': pm_default}.get(slot, allday_default)
    return val or 0


def _slot_formula(rate, slot):
    if not rate:
        return ''
    return {'AM': rate['am_value_formula'], 'PM': rate['pm_value_formula']}.get(
        slot, rate['allday_value_formula']) or ''


def _effective_time_weight(time_weight_override, slot):
    """How much of a 'day' a block represents, in half-day units: AM/PM = 1, All day = 2
    by default. A block can override this (e.g. Freshwater's 7:30am-4pm Scope session is
    longer than a normal AM slot, so it's manually weighted at 2.5) - used anywhere we're
    measuring time/effort spent rather than counting raw slots."""
    if time_weight_override is not None:
        return time_weight_override
    return 2.0 if slot == 'All day' else 1.0


def get_or_create_location(db, name):
    name = (name or '').strip() or 'Unknown'
    row = db.execute('SELECT id FROM locations WHERE LOWER(name)=LOWER(?)', (name,)).fetchone()
    if row:
        return row['id']
    max_sort = db.execute('SELECT COALESCE(MAX(sort_order),0)+1 as m FROM locations').fetchone()['m']
    cur = db.execute('''INSERT INTO locations(name, category, color, sort_order)
                        VALUES(?, 'other', '#888888', ?)''', (name, max_sort))
    db.commit()
    return cur.lastrowid


def get_or_create_rate(db, location_id, activity):
    row = db.execute('SELECT * FROM location_activity_rates WHERE location_id=? AND activity=?',
                      (location_id, activity)).fetchone()
    if row:
        return row
    default_color = SCOPE_COLOR if activity == 'Scope' else ''
    db.execute('INSERT INTO location_activity_rates(location_id, activity, color) VALUES(?,?,?)',
               (location_id, activity, default_color))
    db.commit()
    return db.execute('SELECT * FROM location_activity_rates WHERE location_id=? AND activity=?',
                       (location_id, activity)).fetchone()


def get_or_create_month(db, year, month, source='import'):
    row = db.execute('SELECT id FROM months WHERE year=? AND month=?', (year, month)).fetchone()
    if row:
        return row['id']
    label = datetime.date(year, month, 1).strftime('%B %Y')
    cur = db.execute('INSERT INTO months(year, month, label, status, source) VALUES(?,?,?,?,?)',
                      (year, month, label, 'draft', source))
    db.commit()
    return cur.lastrowid


def upsert_block(db, month_id, row):
    location_id = get_or_create_location(db, row['site_name'])
    get_or_create_rate(db, location_id, row['activity'])
    notes = "Site marked with '*' in source file" if row.get('site_note') else ''
    existing = db.execute(
        "SELECT id FROM blocks WHERE date=? AND slot=? AND status != 'hold'",
        (row['date'], row['slot'])
    ).fetchone()
    if existing:
        db.execute('''UPDATE blocks SET month_id=?, day_name=?, location_id=?, activity=?,
                       activity_raw=?, time_note=?, notes=?, updated_at=CURRENT_TIMESTAMP
                       WHERE id=?''',
                   (month_id, row['day_name'], location_id, row['activity'],
                    row['activity_raw'], row['time_note'], notes, existing['id']))
        return 'updated'
    db.execute('''INSERT INTO blocks(month_id, date, day_name, slot, location_id,
                   activity, activity_raw, time_note, notes, status, source)
                   VALUES(?,?,?,?,?,?,?,?,?,'scheduled','import')''',
               (month_id, row['date'], row['day_name'], row['slot'], location_id,
                row['activity'], row['activity_raw'], row['time_note'], notes))
    return 'inserted'


_BLOCK_COLS = ['id', 'month_id', 'date', 'day_name', 'slot', 'location_id', 'activity',
               'activity_raw', 'time_note', 'value_override', 'time_weight', 'status', 'notes',
               'source', 'created_at', 'updated_at']


def has_pending_changes(db, month_id):
    return db.execute('SELECT 1 FROM board_snapshots WHERE month_id=?', (month_id,)).fetchone() is not None


def ensure_snapshot(db, month_id):
    """Take a snapshot of month_id's blocks the first time it's edited since the last
    save (or since import) - this is the 'original' state Undo reverts back to. A no-op
    if a snapshot already exists (we only ever keep the OLDEST one, not the most recent)."""
    if has_pending_changes(db, month_id):
        return
    rows = db.execute('SELECT * FROM blocks WHERE month_id=?', (month_id,)).fetchall()
    data = [dict(r) for r in rows]
    db.execute('INSERT INTO board_snapshots(month_id, snapshot_json) VALUES(?,?)',
               (month_id, json.dumps(data)))
    db.commit()


def block_month_id(db, block_id):
    row = db.execute('SELECT month_id FROM blocks WHERE id=?', (block_id,)).fetchone()
    return row['month_id'] if row else None


def undo_month(db, month_id):
    row = db.execute('SELECT snapshot_json FROM board_snapshots WHERE month_id=?', (month_id,)).fetchone()
    if not row:
        return False
    snapshot_blocks = json.loads(row['snapshot_json'])
    snapshot_ids = {b['id'] for b in snapshot_blocks}
    current_ids = {r['id'] for r in db.execute('SELECT id FROM blocks WHERE month_id=?', (month_id,)).fetchall()}
    for extra_id in current_ids - snapshot_ids:
        db.execute('DELETE FROM blocks WHERE id=?', (extra_id,))
    for b in snapshot_blocks:
        placeholders = ','.join(f':{c}' for c in _BLOCK_COLS)
        db.execute(f'INSERT OR REPLACE INTO blocks ({",".join(_BLOCK_COLS)}) VALUES ({placeholders})', b)
    db.execute('DELETE FROM board_snapshots WHERE month_id=?', (month_id,))
    db.commit()
    return True


def save_month(db, month_id):
    db.execute('DELETE FROM board_snapshots WHERE month_id=?', (month_id,))
    db.commit()


def get_board_data(db, month_id):
    blocks = db.execute('''
        SELECT b.*, l.name as location_name, l.color as location_color,
               l.default_am_value, l.default_pm_value, l.default_allday_value
        FROM blocks b LEFT JOIN locations l ON l.id = b.location_id
        WHERE b.month_id=? AND b.status != 'hold'
        ORDER BY b.date, CASE b.slot WHEN 'AM' THEN 0 WHEN 'PM' THEN 1 ELSE 2 END
    ''', (month_id,)).fetchall()
    rates = db.execute('SELECT * FROM location_activity_rates').fetchall()
    rate_map = {(r['location_id'], r['activity']): r for r in rates}

    by_date = {}
    for b in blocks:
        rate = rate_map.get((b['location_id'], b['activity']))
        color = _effective_color(rate, b['location_color'])
        value = _effective_value(b['value_override'], rate, b['default_am_value'],
                                  b['default_pm_value'], b['default_allday_value'], b['slot'])
        value_formula = _slot_formula(rate, b['slot'])
        time_weight = _effective_time_weight(b['time_weight'], b['slot'])
        by_date.setdefault(b['date'], []).append(
            dict(b, color=color, value=value, value_formula=value_formula, time_weight_effective=time_weight))

    dates = sorted(by_date.keys())
    if not dates:
        return []
    weeks = []
    current_monday = None
    current_week = None
    for d_str in dates:
        d = datetime.date.fromisoformat(d_str)
        monday = d - datetime.timedelta(days=d.weekday())
        if monday != current_monday:
            current_monday = monday
            current_week = {'monday': monday, 'label': week_rotation_label(monday), 'days': {}}
            weeks.append(current_week)
        day_blocks = by_date[d_str]
        current_week['days'][d.weekday()] = {'date': d_str, 'day_name': d.strftime('%A'), 'blocks': day_blocks,
                                              'total': sum(b['value'] for b in day_blocks)}
    for w in weeks:
        days = []
        for wd in range(5):
            if wd in w['days']:
                days.append(w['days'][wd])
            else:
                d = w['monday'] + datetime.timedelta(days=wd)
                days.append({'date': d.isoformat(), 'day_name': d.strftime('%A'), 'blocks': [], 'total': 0})
        w['days'] = days
        w['total'] = sum(day['total'] for day in days)
    return weeks


def get_hold_blocks(db):
    blocks = db.execute('''
        SELECT b.*, l.name as location_name, l.color as location_color
        FROM blocks b LEFT JOIN locations l ON l.id=b.location_id
        WHERE b.status='hold' ORDER BY b.updated_at DESC
    ''').fetchall()
    rates = db.execute('SELECT * FROM location_activity_rates').fetchall()
    rate_map = {(r['location_id'], r['activity']): r for r in rates}
    out = []
    for b in blocks:
        rate = rate_map.get((b['location_id'], b['activity']))
        color = _effective_color(rate, b['location_color'])
        value_formula = _slot_formula(rate, b['slot'])
        time_weight = _effective_time_weight(b['time_weight'], b['slot'])
        out.append(dict(b, color=color, value_formula=value_formula, time_weight_effective=time_weight))
    return out


def compute_earnings(db, month_id):
    month_row = db.execute('SELECT year, month FROM months WHERE id=?', (month_id,)).fetchone()
    ym_prefix = f"{month_row['year']:04d}-{month_row['month']:02d}" if month_row else ''
    blocks = db.execute('''SELECT b.*, l.name as location_name, l.default_am_value,
                           l.default_pm_value, l.default_allday_value
                           FROM blocks b LEFT JOIN locations l ON l.id=b.location_id
                           WHERE b.month_id=? AND b.status != 'hold' ''', (month_id,)).fetchall()
    rates = db.execute('SELECT * FROM location_activity_rates').fetchall()
    rate_map = {(r['location_id'], r['activity']): r for r in rates}
    weekly, by_site, total = {}, {}, 0.0
    for b in blocks:
        rate = rate_map.get((b['location_id'], b['activity']))
        value = _effective_value(b['value_override'], rate, b['default_am_value'],
                                  b['default_pm_value'], b['default_allday_value'], b['slot'])
        d = datetime.date.fromisoformat(b['date'])
        monday = (d - datetime.timedelta(days=d.weekday())).isoformat()
        # weekly rollup shows the real week total (may include a spillover day from an
        # adjacent month at the boundary); total/by_site only count this month's own days,
        # so a boundary week's earnings aren't double-counted or miscredited between months.
        weekly[monday] = weekly.get(monday, 0.0) + value
        if b['date'].startswith(ym_prefix):
            site_name = b['location_name'] or 'Unknown'
            by_site[site_name] = by_site.get(site_name, 0.0) + value
            total += value
    return {'weekly': sorted(weekly.items()), 'by_site': sorted(by_site.items(), key=lambda x: -x[1]), 'total': total}


def _compute_analysis(db, period):
    """period is either a specific month_id (int) or the literal string 'all' (whole year -
    every non-archived month combined). Returns one dict with everything the Analysis page
    and the Claude narrative both need, or None if there's no data yet for that scope."""
    ym_prefix = None
    if period == 'all':
        months = db.execute("SELECT * FROM months WHERE status != 'archived' ORDER BY year, month").fetchall()
        month_ids = [m['id'] for m in months]
        n_months = len(months) or 1
        period_label = 'Whole year'
    else:
        month = db.execute('SELECT * FROM months WHERE id=?', (period,)).fetchone()
        month_ids = [month['id']] if month else []
        n_months = 1
        period_label = month['label'] if month else 'Unknown month'
        if month:
            ym_prefix = f"{month['year']:04d}-{month['month']:02d}"

    if not month_ids:
        return None

    placeholders = ','.join('?' * len(month_ids))
    rows = db.execute(f'''
        SELECT l.id as location_id, l.name as location_name, l.category as category,
               b.date, b.value_override, b.slot, b.activity, b.time_weight,
               l.default_am_value, l.default_pm_value, l.default_allday_value
        FROM blocks b JOIN locations l ON l.id = b.location_id
        WHERE b.month_id IN ({placeholders}) AND b.status != 'hold'
    ''', month_ids).fetchall()
    if ym_prefix:
        # A single-month view should only count that month's own days - a boundary week's
        # spillover into the next month (e.g. August's board showing early September)
        # belongs in that OTHER month's analysis, not here. "Whole year" has no such
        # exclusion since every day belongs to the year regardless.
        rows = [r for r in rows if r['date'].startswith(ym_prefix)]

    rates = db.execute('SELECT * FROM location_activity_rates').fetchall()
    rate_map = {(r['location_id'], r['activity']): r for r in rates}

    site_counts, site_totals, site_time_units = {}, {}, {}
    site_activity_counts, site_activity_totals = {}, {}
    group_counts, group_totals, group_time_units = {}, {}, {}
    total_sessions, total_revenue, total_time_units = 0, 0.0, 0.0

    # Rotation-week breakdown (Week 1-4) rather than calendar-week-starting dates: the
    # rotation label is stable across months/years, so this is the useful frame for
    # comparing "how does Week 3 perform" over time, unlike shifting calendar dates.
    rotation = {label: {'sessions': 0, 'value': 0.0,
                         'days': {w: {'sessions': 0, 'value': 0.0} for w in WEEKDAY_ORDER}}
                for label in ROTATION_LABELS}

    for r in rows:
        rate = rate_map.get((r['location_id'], r['activity']))
        value = _effective_value(r['value_override'], rate, r['default_am_value'],
                                  r['default_pm_value'], r['default_allday_value'], r['slot'])
        weight = _effective_time_weight(r['time_weight'], r['slot'])
        name = r['location_name']
        site_counts[name] = site_counts.get(name, 0) + 1
        site_totals[name] = site_totals.get(name, 0.0) + value
        site_time_units[name] = site_time_units.get(name, 0.0) + weight
        key = (name, r['activity'])
        site_activity_counts[key] = site_activity_counts.get(key, 0) + 1
        site_activity_totals[key] = site_activity_totals.get(key, 0.0) + value

        group = classify_group(r['category'], r['activity'])
        group_counts[group] = group_counts.get(group, 0) + 1
        group_totals[group] = group_totals.get(group, 0.0) + value
        group_time_units[group] = group_time_units.get(group, 0.0) + weight

        total_sessions += 1
        total_revenue += value
        total_time_units += weight

        d = datetime.date.fromisoformat(r['date'])
        monday = d - datetime.timedelta(days=d.weekday())
        weekday_name = d.strftime('%A')
        if weekday_name in WEEKDAY_ORDER:
            wk = rotation[week_rotation_label(monday)]
            wk['sessions'] += 1
            wk['value'] += value
            wk['days'][weekday_name]['sessions'] += 1
            wk['days'][weekday_name]['value'] += value

    rotation_weeks = [{'label': label, **rotation[label]} for label in ROTATION_LABELS]

    site_breakdown = []
    for site, rev in sorted(site_totals.items(), key=lambda x: -x[1]):
        cnt = site_counts[site]
        tu = site_time_units.get(site, 0.0)
        site_breakdown.append({
            'site': site, 'sessions': cnt, 'revenue': rev,
            'pct': round(100 * rev / total_revenue, 1) if total_revenue else 0.0,
            'avg_per_month': round(cnt / n_months, 2),
            'time_units': tu,
            'revenue_per_time_unit': round(rev / tu, 2) if tu else 0.0,
        })

    group_breakdown = []
    for group in GROUP_ORDER + (['Other'] if group_counts.get('Other') else []):
        cnt = group_counts.get(group, 0)
        rev = group_totals.get(group, 0.0)
        tu = group_time_units.get(group, 0.0)
        group_breakdown.append({
            'group': group, 'sessions': cnt, 'revenue': rev,
            'pct': round(100 * rev / total_revenue, 1) if total_revenue else 0.0,
            'avg_revenue_per_session': round(rev / cnt, 2) if cnt else 0.0,
            'time_units': tu,
            'revenue_per_time_unit': round(rev / tu, 2) if tu else 0.0,
        })

    # Efficiency table: $ earned per half-day-equivalent time unit spent, sites/sites with
    # real time logged only, worst (lowest $/time-unit) first - flags "there for hours,
    # earning little" rather than just low session count (which utilization flags already cover).
    efficiency = sorted(
        (b for b in site_breakdown if b['time_units']),
        key=lambda b: b['revenue_per_time_unit']
    )

    # Scope vs consult time-spent ratio (in half-day-equivalent units, not raw slot counts,
    # so a Dee Why all-day counts as 2 and Freshwater's extended day counts as its real
    # weight) - CDD deliberately excluded (it does both, and isn't part of either referral
    # pipeline below), matching the user's own framing.
    scope_sessions = group_counts.get('Scopes', 0)
    consult_sessions = group_counts.get('FORHEALTH Medical Centre', 0) + group_counts.get('SHORE Gastroenterology', 0)
    scope_time_units = group_time_units.get('Scopes', 0.0)
    consult_time_units = group_time_units.get('FORHEALTH Medical Centre', 0.0) + group_time_units.get('SHORE Gastroenterology', 0.0)
    cdd_time_units = group_time_units.get('CDD', 0.0)
    scope_revenue = group_totals.get('Scopes', 0.0)
    consult_revenue = group_totals.get('FORHEALTH Medical Centre', 0.0) + group_totals.get('SHORE Gastroenterology', 0.0)
    consult_scope_ratio = round(consult_time_units / scope_time_units, 1) if scope_time_units else None

    under_threshold = float(cfg('utilization_flag_threshold', '1') or 1)
    over_threshold = float(cfg('analysis_overbook_threshold', '10') or 10)
    flags = []
    for b in site_breakdown:
        if b['avg_per_month'] <= under_threshold:
            flags.append({'site': b['site'], 'type': 'under-utilized',
                          'detail': f"{b['sessions']} session(s) - avg {b['avg_per_month']}/month"})
        elif b['avg_per_month'] >= over_threshold:
            flags.append({'site': b['site'], 'type': 'over-booked',
                          'detail': f"{b['sessions']} session(s) - avg {b['avg_per_month']}/month"})

    pipelines = []
    for p in REFERRAL_PIPELINES:
        source_sessions = sum(site_counts.get(s, 0) for s in p['source_sites'])
        source_revenue = sum(site_totals.get(s, 0.0) for s in p['source_sites'])
        if p['source_group']:
            source_sessions += group_counts.get(p['source_group'], 0)
            source_revenue += group_totals.get(p['source_group'], 0.0)
        target_sessions = sum(site_counts.get(s, 0) for s in p['target_sites'])
        target_revenue = sum(site_totals.get(s, 0.0) for s in p['target_sites'])
        for key in p['target_site_activity']:
            target_sessions += site_activity_counts.get(key, 0)
            target_revenue += site_activity_totals.get(key, 0.0)
        pipelines.append({
            'label': p['label'],
            'source_sessions': source_sessions, 'source_revenue': source_revenue,
            'target_sessions': target_sessions, 'target_revenue': target_revenue,
            'ratio': round(source_sessions / target_sessions, 1) if target_sessions else None,
        })

    return {
        'period_label': period_label,
        'n_months': n_months,
        'total_sessions': total_sessions,
        'total_revenue': total_revenue,
        'total_time_units': total_time_units,
        'site_breakdown': site_breakdown,
        'group_breakdown': group_breakdown,
        'efficiency': efficiency,
        'flags': flags,
        'rotation_weeks': rotation_weeks,
        'scope_sessions': scope_sessions,
        'consult_sessions': consult_sessions,
        'scope_time_units': scope_time_units,
        'consult_time_units': consult_time_units,
        'cdd_time_units': cdd_time_units,
        'scope_revenue': scope_revenue,
        'consult_revenue': consult_revenue,
        'consult_scope_ratio': consult_scope_ratio,
        'pipelines': pipelines,
    }


# ---------------------------------------------------------------- Board routes

@app.route('/')
def index():
    db = get_db()
    row = db.execute("SELECT id FROM months WHERE status != 'archived' ORDER BY year DESC, month DESC LIMIT 1").fetchone()
    if row:
        return redirect(url_for('board', month_id=row['id']))
    return redirect(url_for('import_schedule'))


@app.route('/board/<int:month_id>')
def board(month_id):
    db = get_db()
    month = db.execute('SELECT * FROM months WHERE id=?', (month_id,)).fetchone()
    if not month:
        flash('Month not found.', 'danger')
        return redirect(url_for('index'))
    weeks = get_board_data(db, month_id)
    holds = get_hold_blocks(db)
    all_months = db.execute("SELECT * FROM months WHERE status != 'archived' ORDER BY year, month").fetchall()
    locations = db.execute('SELECT * FROM locations WHERE active=1 ORDER BY sort_order, name').fetchall()
    pending = has_pending_changes(db, month_id)
    # Only count days that actually fall within this calendar month - a boundary week's
    # spillover days (e.g. August's board showing into early September) belong to that
    # OTHER month's total, not this one, even though they're displayed here for context.
    ym_prefix = f"{month['year']:04d}-{month['month']:02d}"
    month_total = sum(day['total'] for w in weeks for day in w['days'] if day['date'].startswith(ym_prefix))
    return render_template('board.html', month=month, weeks=weeks, holds=holds,
                           all_months=all_months, locations=locations, activities=ACTIVITIES,
                           pending=pending, month_total=month_total)


@app.route('/board/<int:month_id>/undo', methods=['POST'])
def undo_board(month_id):
    db = get_db()
    if undo_month(db, month_id):
        flash('Changes undone - back to the last saved state.', 'success')
    else:
        flash('Nothing to undo.', 'info')
    return redirect(url_for('board', month_id=month_id))


@app.route('/board/<int:month_id>/save', methods=['POST'])
def save_board(month_id):
    db = get_db()
    save_month(db, month_id)
    flash('Changes saved.', 'success')
    return redirect(url_for('board', month_id=month_id))


@app.route('/block/<int:block_id>/move', methods=['POST'])
def move_block(block_id):
    data = request.get_json(silent=True) or {}
    db = get_db()
    month_id = block_month_id(db, block_id)
    if month_id:
        ensure_snapshot(db, month_id)
    if data.get('status') == 'hold':
        db.execute("UPDATE blocks SET status='hold', updated_at=CURRENT_TIMESTAMP WHERE id=?", (block_id,))
    else:
        db.execute('''UPDATE blocks SET date=?, slot=?, status='scheduled', updated_at=CURRENT_TIMESTAMP
                       WHERE id=?''', (data.get('date'), data.get('slot'), block_id))
    db.commit()
    return jsonify({'ok': True})


@app.route('/block/<int:block_id>/set-value', methods=['POST'])
def set_block_value(block_id):
    data = request.get_json(silent=True) or {}
    value = data.get('value_override', None)
    db = get_db()
    if value in (None, ''):
        db.execute('UPDATE blocks SET value_override=NULL, updated_at=CURRENT_TIMESTAMP WHERE id=?', (block_id,))
    else:
        try:
            value = float(value)
        except (TypeError, ValueError):
            return jsonify({'ok': False, 'error': 'Invalid value'}), 400
        db.execute('UPDATE blocks SET value_override=?, updated_at=CURRENT_TIMESTAMP WHERE id=?', (value, block_id))
    db.commit()
    return jsonify({'ok': True})


@app.route('/block/<int:block_id>/set-notes', methods=['POST'])
def set_block_notes(block_id):
    data = request.get_json(silent=True) or {}
    db = get_db()
    month_id = block_month_id(db, block_id)
    if month_id:
        ensure_snapshot(db, month_id)
    db.execute('UPDATE blocks SET notes=?, updated_at=CURRENT_TIMESTAMP WHERE id=?',
               (data.get('notes', ''), block_id))
    db.commit()
    return jsonify({'ok': True})


@app.route('/block/<int:block_id>/set-time-weight', methods=['POST'])
def set_block_time_weight(block_id):
    """How much of a 'day' THIS specific session represents, in half-day units (AM/PM
    default to 1, All day defaults to 2) - for sessions that run longer or shorter than
    the norm (e.g. a Scope that runs 7:30am-4pm), so time/effort-based analysis reflects
    reality rather than just counting slots. Blank clears the override back to the default."""
    data = request.get_json(silent=True) or {}
    raw = data.get('time_weight', None)
    db = get_db()
    month_id = block_month_id(db, block_id)
    if month_id:
        ensure_snapshot(db, month_id)
    if raw in (None, ''):
        db.execute('UPDATE blocks SET time_weight=NULL, updated_at=CURRENT_TIMESTAMP WHERE id=?', (block_id,))
    else:
        try:
            weight = float(raw)
        except (TypeError, ValueError):
            return jsonify({'ok': False, 'error': 'Time weight must be a number'}), 400
        db.execute('UPDATE blocks SET time_weight=?, updated_at=CURRENT_TIMESTAMP WHERE id=?', (weight, block_id))
    db.commit()
    return jsonify({'ok': True})


@app.route('/block/<int:block_id>/set-activity', methods=['POST'])
def set_block_activity(block_id):
    """Rename just THIS block's activity label (e.g. 'Consult' -> 'Evening Consult').
    Since colour/value are keyed on (site, activity), giving one session its own unique
    activity name splits it into its own priceable category without touching every other
    session at that site - the opposite of the site-wide value edit above."""
    data = request.get_json(silent=True) or {}
    new_activity = (data.get('activity') or '').strip()
    if not new_activity:
        return jsonify({'ok': False, 'error': 'Activity name cannot be blank'}), 400
    db = get_db()
    block = db.execute('SELECT location_id, month_id FROM blocks WHERE id=?', (block_id,)).fetchone()
    if not block:
        return jsonify({'ok': False, 'error': 'Block not found'}), 404
    if block['month_id']:
        ensure_snapshot(db, block['month_id'])
    get_or_create_rate(db, block['location_id'], new_activity)
    db.execute('UPDATE blocks SET activity=?, updated_at=CURRENT_TIMESTAMP WHERE id=?',
               (new_activity, block_id))
    db.commit()
    return jsonify({'ok': True})


@app.route('/block/<int:block_id>/delete', methods=['POST'])
def delete_block(block_id):
    db = get_db()
    month_id = block_month_id(db, block_id)
    if month_id:
        ensure_snapshot(db, month_id)
    db.execute('DELETE FROM blocks WHERE id=?', (block_id,))
    db.commit()
    return jsonify({'ok': True})


@app.route('/block/new', methods=['POST'])
def new_block():
    data = request.get_json(silent=True) or request.form
    db = get_db()
    month_id = data.get('month_id')
    date_str = data.get('date')
    slot = data.get('slot')
    location_id = data.get('location_id')
    activity = data.get('activity')
    if not (date_str and slot and location_id and activity):
        return jsonify({'ok': False, 'error': 'Missing fields'}), 400
    if month_id:
        ensure_snapshot(db, int(month_id))
    location_id = int(location_id)
    get_or_create_rate(db, location_id, activity)
    day_name = datetime.date.fromisoformat(date_str).strftime('%A')
    cur = db.execute('''INSERT INTO blocks(month_id, date, day_name, slot, location_id, activity, status, source)
                        VALUES(?,?,?,?,?,?, 'scheduled', 'manual')''',
                      (month_id, date_str, day_name, slot, location_id, activity))
    db.commit()
    return jsonify({'ok': True, 'id': cur.lastrowid})


# ---------------------------------------------------------------- Import

@app.route('/import', methods=['GET', 'POST'])
def import_schedule():
    db = get_db()
    if request.method == 'POST':
        file = request.files.get('file')
        if not file or not file.filename:
            flash('Please choose a file to upload.', 'warning')
            return redirect(url_for('import_schedule'))
        try:
            wb = openpyxl.load_workbook(BytesIO(file.read()), data_only=True)
        except Exception as e:
            flash(f'Could not read that file: {e}', 'danger')
            return redirect(url_for('import_schedule'))
        try:
            year, month, rows = importer.parse_schedule_list(wb)
        except ValueError as list_err:
            try:
                year, month, rows = importer.parse_calendar_grid(wb)
            except ValueError:
                flash(str(list_err), 'danger')
                return redirect(url_for('import_schedule'))

        month_id = get_or_create_month(db, year, month)
        inserted = updated = 0
        for row in rows:
            result = upsert_block(db, month_id, row)
            if result == 'inserted':
                inserted += 1
            else:
                updated += 1
        db.commit()
        db.execute('''INSERT INTO import_log(original_filename, month_id, rows_inserted, rows_updated, rows_skipped)
                       VALUES(?,?,?,?,?)''', (file.filename, month_id, inserted, updated, 0))
        db.commit()
        month_row = db.execute('SELECT label FROM months WHERE id=?', (month_id,)).fetchone()
        flash(f"Imported {len(rows)} rows: {inserted} new, {updated} updated. Month: {month_row['label']}.", 'success')
        return redirect(url_for('board', month_id=month_id))

    logs = db.execute('''SELECT il.*, m.label as month_label FROM import_log il
                          LEFT JOIN months m ON m.id = il.month_id
                          ORDER BY il.imported_at DESC LIMIT 20''').fetchall()
    return render_template('import.html', logs=logs)


# ---------------------------------------------------------------- Locations / Settings

@app.route('/locations')
def locations_page():
    db = get_db()
    locations = db.execute('SELECT * FROM locations ORDER BY sort_order, name').fetchall()
    rates = db.execute('''SELECT r.*, l.name as location_name FROM location_activity_rates r
                           JOIN locations l ON l.id = r.location_id
                           ORDER BY l.sort_order, l.name, r.activity''').fetchall()
    return render_template('locations.html', locations=locations, rates=rates, categories=CATEGORIES)


@app.route('/locations/<int:loc_id>/update', methods=['POST'])
def update_location(loc_id):
    f = request.form
    db = get_db()
    db.execute('''UPDATE locations SET name=?, category=?, color=?, default_am_value=?,
                   default_pm_value=?, default_allday_value=?, active=? WHERE id=?''',
               (f.get('name'), f.get('category'), f.get('color'),
                float(f.get('default_am_value') or 0), float(f.get('default_pm_value') or 0),
                float(f.get('default_allday_value') or 0), 1 if f.get('active') else 0, loc_id))
    db.commit()
    flash('Location updated.', 'success')
    return redirect(url_for('locations_page'))


@app.route('/location-rate/<int:rate_id>/update', methods=['POST'])
def update_location_rate(rate_id):
    f = request.form
    db = get_db()
    try:
        am_value, am_formula = parse_money_input(f.get('am_value'))
        pm_value, pm_formula = parse_money_input(f.get('pm_value'))
        allday_value, allday_formula = parse_money_input(f.get('allday_value'))
    except ValueError as e:
        flash(str(e), 'danger')
        return redirect(url_for('locations_page'))
    db.execute('''UPDATE location_activity_rates SET color=?, am_value=?, pm_value=?, allday_value=?,
                   am_value_formula=?, pm_value_formula=?, allday_value_formula=? WHERE id=?''',
               (f.get('color') or '', am_value, pm_value, allday_value,
                am_formula, pm_formula, allday_formula, rate_id))
    db.commit()
    flash('Rate updated.', 'success')
    return redirect(url_for('locations_page'))


@app.route('/location-rate/set-value', methods=['POST'])
def set_rate_value():
    """Set a shared value for a (location, activity) combo, applied to every session at
    that site/activity across every month - not a per-block override. This is the board's
    quick-edit action; use /locations for finer per-slot control instead.

    All-day sessions are priced separately from half-day (AM/PM) ones - a whole-day
    booking is a different category of work, not just "two half-days back to back" -
    so 'slot' picks which of the two this edit applies to: AM/PM share one value,
    All day has its own."""
    data = request.get_json(silent=True) or {}
    location_id = data.get('location_id')
    activity = data.get('activity')
    slot = data.get('slot', 'AM')
    db = get_db()
    try:
        value, formula = parse_money_input(data.get('value', ''))
    except ValueError as e:
        return jsonify({'ok': False, 'error': str(e)}), 400
    if not location_id or not activity:
        return jsonify({'ok': False, 'error': 'Missing location/activity'}), 400
    get_or_create_rate(db, location_id, activity)
    if slot == 'All day':
        db.execute('''UPDATE location_activity_rates SET allday_value=?, allday_value_formula=?
                       WHERE location_id=? AND activity=?''',
                   (value, formula, location_id, activity))
    else:
        db.execute('''UPDATE location_activity_rates
                       SET am_value=?, pm_value=?, am_value_formula=?, pm_value_formula=?
                       WHERE location_id=? AND activity=?''',
                   (value, value, formula, formula, location_id, activity))
    db.commit()
    return jsonify({'ok': True, 'value': value})


@app.route('/settings', methods=['GET', 'POST'])
def settings_page():
    if request.method == 'POST':
        set_cfg('claude_api_key', request.form.get('claude_api_key', '').strip())
        set_cfg('analysis_trailing_months', request.form.get('analysis_trailing_months', '3'))
        set_cfg('utilization_flag_threshold', request.form.get('utilization_flag_threshold', '1'))
        set_cfg('analysis_overbook_threshold', request.form.get('analysis_overbook_threshold', '10'))
        flash('Settings saved.', 'success')
        return redirect(url_for('settings_page'))
    return render_template('settings.html',
                           claude_api_key=cfg('claude_api_key'),
                           analysis_trailing_months=cfg('analysis_trailing_months', '3'),
                           utilization_flag_threshold=cfg('utilization_flag_threshold', '1'),
                           analysis_overbook_threshold=cfg('analysis_overbook_threshold', '10'))


@app.route('/settings/test-claude', methods=['POST'])
def test_claude():
    api_key = cfg('claude_api_key')
    if not api_key:
        flash('Enter and save a Claude API key first.', 'warning')
        return redirect(url_for('settings_page'))
    try:
        claude_integration.test_connection(api_key)
        flash('Claude connection succeeded.', 'success')
    except Exception as e:
        flash(f'Claude connection failed: {e}', 'danger')
    return redirect(url_for('settings_page'))


# ---------------------------------------------------------------- Export / Import (full data)

EXPORT_TABLES = ['locations', 'location_activity_rates', 'months', 'blocks']
# claude_api_key is deliberately excluded - each deployment manages its own key.
EXPORT_CONFIG_KEYS = ['analysis_trailing_months', 'utilization_flag_threshold', 'analysis_overbook_threshold']


@app.route('/export')
def export_data():
    """Dump every table that defines 'what the schedule looks like' - sites, colours,
    per-slot values/formulas, months, and every block (date/slot/activity/notes/status) -
    as one JSON file. Pairs with /import-data to move this exactly onto another deployment
    (e.g. local -> Railway) without re-entering colours/rates/holds by hand."""
    db = get_db()
    data = {'exported_at': datetime.datetime.now().isoformat(), 'version': 1}
    for table in EXPORT_TABLES:
        rows = db.execute(f'SELECT * FROM {table}').fetchall()
        data[table] = [dict(r) for r in rows]
    cfg_rows = db.execute(
        f"SELECT key, value FROM config WHERE key IN ({','.join('?' * len(EXPORT_CONFIG_KEYS))})",
        EXPORT_CONFIG_KEYS).fetchall()
    data['config'] = {r['key']: r['value'] for r in cfg_rows}
    payload = json.dumps(data, indent=2)
    filename = f"practice_manager_export_{datetime.date.today().isoformat()}.json"
    return Response(payload, mimetype='application/json',
                     headers={'Content-Disposition': f'attachment; filename="{filename}"'})


@app.route('/import-data', methods=['GET', 'POST'])
def import_data_page():
    if request.method == 'POST':
        file = request.files.get('file')
        if not file or not file.filename:
            flash('Please choose an export file to upload.', 'warning')
            return redirect(url_for('import_data_page'))
        try:
            payload = json.loads(file.read().decode('utf-8'))
        except (ValueError, UnicodeDecodeError) as e:
            flash(f'Could not read that file: {e}', 'danger')
            return redirect(url_for('import_data_page'))

        db = get_db()
        counts = {t: len(payload.get(t, [])) for t in EXPORT_TABLES}
        try:
            # Child tables first, then parents, mirroring foreign-key direction both ways.
            # import_log.month_id also references months (no cascade) - it's just an audit
            # trail, not schedule data, so it's cleared rather than exported/imported.
            db.execute('DELETE FROM blocks')
            db.execute('DELETE FROM location_activity_rates')
            db.execute('DELETE FROM board_snapshots')
            db.execute('DELETE FROM import_log')
            db.execute('DELETE FROM months')
            db.execute('DELETE FROM locations')

            for loc in payload.get('locations', []):
                cols = list(loc.keys())
                db.execute(f'INSERT INTO locations ({",".join(cols)}) VALUES ({",".join(":" + c for c in cols)})', loc)
            for m in payload.get('months', []):
                cols = list(m.keys())
                db.execute(f'INSERT INTO months ({",".join(cols)}) VALUES ({",".join(":" + c for c in cols)})', m)
            for rate in payload.get('location_activity_rates', []):
                cols = list(rate.keys())
                db.execute(f'INSERT INTO location_activity_rates ({",".join(cols)}) '
                           f'VALUES ({",".join(":" + c for c in cols)})', rate)
            for b in payload.get('blocks', []):
                cols = list(b.keys())
                db.execute(f'INSERT INTO blocks ({",".join(cols)}) VALUES ({",".join(":" + c for c in cols)})', b)
            for k, v in payload.get('config', {}).items():
                if k in EXPORT_CONFIG_KEYS:
                    db.execute('INSERT OR REPLACE INTO config(key, value) VALUES (?, ?)', (k, v))
            db.commit()
        except Exception as e:
            db.rollback()
            flash(f'Import failed, no changes were made: {e}', 'danger')
            return redirect(url_for('import_data_page'))

        flash(f"Import complete - replaced everything with: {counts['locations']} locations, "
              f"{counts['location_activity_rates']} site/activity rates, {counts['months']} months, "
              f"{counts['blocks']} blocks.", 'success')
        return redirect(url_for('index'))

    return render_template('import_data.html')


# ---------------------------------------------------------------- Months

@app.route('/months')
def months_page():
    db = get_db()
    months = db.execute("SELECT * FROM months WHERE status != 'archived' ORDER BY year DESC, month DESC").fetchall()
    return render_template('months.html', months=months, archived=False)


@app.route('/months/archive')
def months_archive_page():
    db = get_db()
    months = db.execute("SELECT * FROM months WHERE status = 'archived' ORDER BY year DESC, month DESC").fetchall()
    return render_template('months.html', months=months, archived=True)


@app.route('/months/<int:month_id>/publish', methods=['POST'])
def publish_month(month_id):
    db = get_db()
    holds = db.execute("SELECT COUNT(*) as c FROM blocks WHERE month_id=? AND status='hold'",
                        (month_id,)).fetchone()['c']
    db.execute("UPDATE months SET status='published', published_at=CURRENT_TIMESTAMP WHERE id=?", (month_id,))
    save_month(db, month_id)
    db.commit()
    msg = 'Month published.'
    if holds:
        msg += f' Note: {holds} block(s) are still on hold and will not appear as scheduled.'
    flash(msg, 'success')
    return redirect(url_for('months_page'))


@app.route('/months/<int:month_id>/archive', methods=['POST'])
def archive_month(month_id):
    db = get_db()
    db.execute("UPDATE months SET status='archived', archived_at=CURRENT_TIMESTAMP WHERE id=?", (month_id,))
    db.commit()
    flash('Month archived.', 'success')
    return redirect(url_for('months_page'))


@app.route('/months/<int:month_id>/unarchive', methods=['POST'])
def unarchive_month(month_id):
    db = get_db()
    db.execute("UPDATE months SET status='draft', archived_at=NULL WHERE id=?", (month_id,))
    db.commit()
    flash('Month unarchived.', 'success')
    return redirect(url_for('months_archive_page'))


# ---------------------------------------------------------------- Earnings

@app.route('/earnings')
def earnings_page():
    db = get_db()
    month_id = request.args.get('month_id', type=int)
    months = db.execute("SELECT * FROM months WHERE status != 'archived' ORDER BY year DESC, month DESC").fetchall()
    if not month_id and months:
        month_id = months[0]['id']
    data = compute_earnings(db, month_id) if month_id else None
    month = db.execute('SELECT * FROM months WHERE id=?', (month_id,)).fetchone() if month_id else None
    is_actual = bool(month and month['status'] == 'published')
    return render_template('earnings.html', months=months, month_id=month_id, month=month,
                           data=data, is_actual=is_actual)


# ---------------------------------------------------------------- Analysis

def _resolve_period(raw, months):
    if raw == 'all':
        return 'all'
    if raw and raw.isdigit():
        return int(raw)
    return months[0]['id'] if months else 'all'


@app.route('/analysis')
def analysis_page():
    db = get_db()
    months = db.execute("SELECT * FROM months WHERE status != 'archived' ORDER BY year DESC, month DESC").fetchall()
    period = _resolve_period(request.args.get('period'), months)
    data = _compute_analysis(db, period)
    return render_template('analysis.html', data=data, months=months, period=period, narrative=None)


@app.route('/analysis/claude', methods=['POST'])
def analysis_claude():
    db = get_db()
    months = db.execute("SELECT * FROM months WHERE status != 'archived' ORDER BY year DESC, month DESC").fetchall()
    period = _resolve_period(request.form.get('period'), months)
    data = _compute_analysis(db, period)
    api_key = cfg('claude_api_key')
    narrative = None
    if not api_key:
        flash('Set your Claude API key in Settings first.', 'warning')
    elif not data:
        flash('No data for this period yet.', 'warning')
    else:
        try:
            narrative = claude_integration.narrative_analysis(api_key, data)
        except Exception as e:
            flash(f'Claude request failed: {e}', 'danger')
    return render_template('analysis.html', data=data, months=months, period=period, narrative=narrative)


# ---------------------------------------------------------------- Claude schedule generation

@app.route('/generate', methods=['GET', 'POST'])
def generate_page():
    db = get_db()
    months = db.execute("SELECT * FROM months ORDER BY year DESC, month DESC").fetchall()
    if request.method == 'POST':
        api_key = cfg('claude_api_key')
        if not api_key:
            flash('Set your Claude API key in Settings first.', 'warning')
            return redirect(url_for('generate_page'))
        target_year = request.form.get('target_year', type=int)
        target_month = request.form.get('target_month', type=int)
        reference_month_id = request.form.get('reference_month_id', type=int)
        constraints_text = request.form.get('constraints', '')

        existing = db.execute('SELECT id FROM months WHERE year=? AND month=?',
                               (target_year, target_month)).fetchone()
        if existing:
            flash(f'{target_year}-{target_month:02d} already has a month record — pick a month with '
                  f'no existing schedule, or edit that month directly on the board.', 'warning')
            return redirect(url_for('generate_page'))

        ref_blocks = db.execute('''SELECT b.date, b.day_name, b.slot, l.name as site, b.activity, b.notes
                                    FROM blocks b JOIN locations l ON l.id=b.location_id
                                    WHERE b.month_id=? AND b.status != 'hold' ORDER BY b.date''',
                                 (reference_month_id,)).fetchall()
        ref_blocks = [dict(r) for r in ref_blocks]
        location_names = [r['name'] for r in db.execute('SELECT name FROM locations WHERE active=1').fetchall()]

        try:
            draft_rows = claude_integration.generate_schedule_draft(
                api_key, ref_blocks, location_names, target_year, target_month, constraints_text)
        except Exception as e:
            flash(f'Claude schedule generation failed: {e}', 'danger')
            return redirect(url_for('generate_page'))

        label = datetime.date(target_year, target_month, 1).strftime('%B %Y')
        cur = db.execute('''INSERT INTO months(year, month, label, status, source)
                            VALUES(?,?,?,'draft','claude_generated')''', (target_year, target_month, label))
        db.commit()
        month_id = cur.lastrowid

        created = 0
        for r in draft_rows:
            site_name = (r.get('site') or '').strip()
            if not site_name:
                continue
            loc = db.execute('SELECT id FROM locations WHERE LOWER(name)=LOWER(?)', (site_name,)).fetchone()
            if loc:
                location_id = loc['id']
            else:
                cur2 = db.execute('''INSERT INTO locations(name, category, color, active)
                                     VALUES(?, 'other', '#888888', 0)''', (site_name,))
                db.commit()
                location_id = cur2.lastrowid
            activity = r.get('activity') or 'Other'
            get_or_create_rate(db, location_id, activity)
            date_str = r.get('date')
            slot = r.get('slot') or 'AM'
            try:
                day_name = datetime.date.fromisoformat(date_str).strftime('%A')
            except (TypeError, ValueError):
                continue
            db.execute('''INSERT INTO blocks(month_id, date, day_name, slot, location_id, activity, notes,
                          status, source) VALUES(?,?,?,?,?,?,?, 'scheduled', 'claude_generated')''',
                       (month_id, date_str, day_name, slot, location_id, activity, r.get('note') or ''))
            created += 1
        db.commit()
        flash(f'Claude drafted {created} sessions for {label}. Review and edit on the board, then Publish '
              f'when ready.', 'success')
        return redirect(url_for('board', month_id=month_id))

    return render_template('generate.html', months=months)


if __name__ == '__main__':
    init_db()
    app.run(debug=False, host='0.0.0.0', port=int(os.environ.get('PORT', 5003)))
